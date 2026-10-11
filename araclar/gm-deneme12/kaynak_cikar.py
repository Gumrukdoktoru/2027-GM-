"""Kaynak metinlerini metin/ altına çıkarır. Kullanım: python3 kaynak_cikar.py"""
import sys, re, pathlib, subprocess, shutil, hashlib, tempfile
sys.path.insert(0, '/home/claude/2027-ye-haz-rl-k/araclar')
import importlib.util
spec = importlib.util.spec_from_file_location('mc', '/home/claude/2027-ye-haz-rl-k/araclar/metin_cikar.py')
src = open('/home/claude/2027-ye-haz-rl-k/araclar/metin_cikar.py').read().split('for f in sorted')[0]
ns = {}; exec(src, ns)
docx_text, html_text = ns['docx_text'], ns['html_text']
OUT = pathlib.Path(__file__).parent / 'metin'; OUT.mkdir(exist_ok=True)
GM = pathlib.Path('/home/claude/2027-gm-'); GMY = pathlib.Path('/home/claude/2027-ye-haz-rl-k')
def clean(t):
    t = re.sub(r'[ \t\xa0]+\n', '\n', t.replace('\r', ''))
    return re.sub(r'\n{3,}', '\n\n', t).strip() + '\n'
def ole_text(p):
    with tempfile.TemporaryDirectory() as d:
        subprocess.run(['soffice', '--headless', '--convert-to', 'txt:Text (encoded):UTF8', '--outdir', d, str(p)], capture_output=True, timeout=180)
        f = pathlib.Path(d) / (p.stem + '.txt')
        return f.read_text(encoding='utf-8', errors='replace') if f.exists() else ''
def pdf_text(p):
    return subprocess.run(['pdftotext', str(p), '-'], capture_output=True, text=True).stdout
def pptx_text(p):
    from pptx import Presentation
    pr = Presentation(str(p)); out = []
    for i, s in enumerate(pr.slides, 1):
        out.append(f'--- Slayt {i} ---')
        for sh in s.shapes:
            if sh.has_text_frame: out.append(sh.text_frame.text)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows: out.append(' | '.join(c.text for c in r.cells))
    return '\n'.join(out)
def text(p):
    suf = p.suffix.lower()
    if suf == '.docx': return docx_text(p)
    if suf == '.pdf': return pdf_text(p)
    if suf == '.pptx': return pptx_text(p)
    if suf == '.doc':
        head = p.read_bytes()[:8]
        return ole_text(p) if head.startswith(b'\xd0\xcf\x11\xe0') else html_text(p)
    return None
SKIP = {'Canli_7_24_Genel_Kultur_Kitabi.pdf'}
seen = {}
for root in (GM, GMY):
    for p in sorted(root.iterdir()):
        if p.is_dir() or p.name in SKIP or p.suffix.lower() not in ('.doc', '.docx', '.pdf', '.pptx'): continue
        if re.match(r'20\d\d-gmy-sinavi|2025_GÜMRÜK_MÜŞAVİR', p.name): continue
        h = hashlib.md5(p.read_bytes()).hexdigest()
        if h in seen: continue
        name = p.stem
        if (OUT / (name + '.txt')).exists() and root == GMY: name = name + '_gmy'
        t = text(p)
        if not t or len(t.strip()) < 50: print('BOŞ/KISA:', p.name); 
        seen[h] = name
        (OUT / (name + '.txt')).write_text(clean(t or ''), encoding='utf-8')
print('yazılan', len(seen))
